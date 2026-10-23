"""Tests for the patent XLSX export: the export_patents_xlsx command, the workbook
builder and the two admin actions. Spec: specs/patent-export.md."""

import datetime
import os
import tempfile
from io import BytesIO

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.contrib.sites.models import Site
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import connection
from django.test import Client, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from organizations.models import Organization, OrganizationUser

from gregory.exports import patents_xlsx
from gregory.exports.patents_xlsx import (
	PUBLICATION_COLS,
	TRIAL_LINK_COLS,
	ExportScope,
	build_patents_workbook,
)
from gregory.exports.xlsx_common import attribution_for_site
from gregory.management.commands.detect_patent_trial_links import trial_window
from gregory.models import (
	OrganizationSite,
	PatentApplicant,
	PatentCategoryAssignment,
	PatentPublication,
	PatentTrialLink,
	Patents,
	Sources,
	Sponsor,
	Subject,
	Team,
	TeamCategory,
	Trials,
)
from sitesettings.models import CustomSetting

User = get_user_model()
D = datetime.date

SHEET_ORDER = [
	"About this file",
	"Multiple Sclerosis",
	"Parkinson",
	"Trial links",
	"Publications",
	"Categories",
	"Glossary",
	"Data source",
]


def rows_by_header(ws, header_row=1):
	headers = [c.value for c in ws[header_row]]
	return [
		dict(zip(headers, [c.value for c in row]))
		for row in ws.iter_rows(min_row=header_row + 1)
		if any(c.value not in (None, "") for c in row)
	]


def headers_of(ws, header_row=1):
	return [c.value for c in ws[header_row]]


def as_date(value):
	return value.date() if isinstance(value, datetime.datetime) else value


def about_values(wb):
	ws = wb["About this file"]
	return {r[0].value: r[1].value for r in ws.iter_rows() if r[0].value}


class PatentExportTestCase(TestCase):
	def setUp(self):
		self.org = Organization.objects.create(name="Org A", slug="org-a")
		self.team = Team.objects.create(name="Team A", organization=self.org, slug="team-a")
		self.site = Site.objects.create(domain="a.example", name="Site A")
		OrganizationSite.objects.create(organization=self.org, site=self.site, is_default=True)
		self.custom = CustomSetting.objects.create(
			site=self.site, title="Project A", admin_email="a@a.example"
		)
		self.ms = Subject.objects.create(
			subject_name="Multiple Sclerosis", subject_slug="ms", team=self.team
		)
		self.pd = Subject.objects.create(
			subject_name="Parkinson", subject_slug="pd", team=self.team
		)
		self.source = Sources.objects.create(
			name="EPO | MS", source_for="patents", team=self.team
		)
		self.novartis = Sponsor.objects.create(
			name="Novartis", slug="novartis", sponsor_type="industry", sponsor_type_source="curated"
		)
		self.cat_auto = TeamCategory.objects.create(
			team=self.team, category_name="Ocrelizumab", category_terms=["ocrelizumab"]
		)
		self.cat_auto.subjects.add(self.ms)
		self.cat_manual = TeamCategory.objects.create(
			team=self.team, category_name="Anti-CD20", category_terms=[]
		)
		self.cat_manual.subjects.add(self.ms)

		self.p1 = self.make_patent(
			"100",
			"Anti-CD20 antibody",
			D(2022, 5, 1),
			subjects=[self.ms],
			claims="1. An antibody.",
			summary="An abstract.",
			inventors=["Jane Doe", "John Roe"],
			has_grant=True,
			representative="EP4123456A1",
			cpc=["A61K39/395", "A61P25/00"],
		)
		self.pub_a1 = self.add_pub(self.p1, "EP", "4123456", "A1", D(2023, 11, 1), "EP22000001", D(2022, 5, 1))
		self.pub_b1 = self.add_pub(self.p1, "EP", "4123456", "B1", D(2025, 2, 1), "EP22000001", D(2022, 5, 1))
		self.pub_wo = self.add_pub(self.p1, "WO", "2022123456", "A1", D(2022, 12, 1), "PCT/EP2022/1", D(2022, 5, 2))
		PatentApplicant.objects.create(
			patent=self.p1, sponsor=self.novartis, raw_name="NOVARTIS AG", sequence=1
		)
		PatentApplicant.objects.create(
			patent=self.p1, raw_name="SMITH, John", sequence=2, is_individual=True
		)
		PatentApplicant.objects.create(
			patent=self.p1, raw_name="ACME HOLDINGS IP", sequence=3
		)
		PatentCategoryAssignment.objects.create(
			patents=self.p1, teamcategory=self.cat_auto, source="automatic"
		)
		PatentCategoryAssignment.objects.create(
			patents=self.p1, teamcategory=self.cat_manual, source="manual"
		)

		self.trial_closed = self.make_trial("Closed trial", "completed", D(2020, 1, 1), D(2022, 1, 1))
		self.trial_open = self.make_trial("Open trial", "recruiting", D(2021, 1, 1), None)
		self.trial_subject = self.make_trial("Subject trial", "completed", D(2020, 1, 1), D(2023, 1, 1))
		self.trial_dismissed = self.make_trial("Dismissed trial", "completed", D(2020, 1, 1), D(2023, 1, 1))
		self.link_cat_during = self.make_link(self.p1, self.trial_open, "during", "category")
		self.link_cat_after = self.make_link(self.p1, self.trial_closed, "after", "category", days=120)
		self.link_subj = self.make_link(self.p1, self.trial_subject, "during", "subject")
		self.link_suppressed = self.make_link(
			self.p1, self.trial_dismissed, "during", "subject", suppressed=True
		)
		self.trial_open.categories.add(self.cat_auto)

		self.p2 = self.make_patent(
			"200", "Newer US patent", D(2023, 1, 1), subjects=[self.ms, self.pd]
		)
		self.add_pub(self.p2, "US", "2023000001", "A1", D(2024, 7, 1))
		self.p3 = self.make_patent("300", "Undated patent", None, subjects=[self.ms])
		self.add_pub(self.p3, "WO", "2021000001", "A1", D(2021, 6, 1))

	# --- factories ---------------------------------------------------------------

	def make_patent(self, family, title, priority, subjects=(), claims="", summary="",
					inventors=None, has_grant=False, representative="", cpc=None):
		patent = Patents.objects.create(
			family_id=family,
			title=title,
			summary=summary,
			claims=claims,
			earliest_priority_date=priority,
			inventors=inventors or [],
			has_grant=has_grant,
			representative_publication=representative,
			cpc_classes=cpc or [],
			link="https://worldwide.espacenet.com/x",
			links={"espacenet": "https://worldwide.espacenet.com/x"},
		)
		patent.subjects.add(*subjects)
		patent.teams.add(self.team)
		patent.sources.add(self.source)
		return patent

	def add_pub(self, patent, country, number, kind, pub_date, app_no=None, app_date=None):
		return PatentPublication.objects.create(
			patent=patent,
			publication_number=f"{country}{number}{kind}",
			country=country,
			doc_number=number,
			kind=kind,
			publication_date=pub_date,
			application_number=app_no,
			application_date=app_date,
			sources=["epo_ops"],
		)

	def make_trial(self, title, status, start, end, subject=None):
		trial = Trials.objects.create(
			title=title,
			link=f"https://example.com/{title.replace(' ', '-')}",
			recruitment_status_normalized=status,
			date_enrollement=start,
			primary_completion_date=end,
			primary_outcome="Relapse rate",
			secondary_outcome="Disability",
		)
		Trials.objects.filter(pk=trial.pk).update(primary_sponsor_normalized=self.novartis)
		trial.subjects.add(subject or self.ms)
		# save() derives the normalised status from the raw one; pin it for the test
		Trials.objects.filter(pk=trial.pk).update(recruitment_status_normalized=status)
		trial.refresh_from_db()
		return trial

	def make_link(self, patent, trial, timing, basis, days=None, suppressed=False, sponsor=None):
		return PatentTrialLink.objects.create(
			patent=patent,
			trial=trial,
			sponsor=sponsor or self.novartis,
			timing=timing,
			basis=basis,
			days_after_completion=days,
			suppressed=suppressed,
		)

	def build(self, subjects=None, **kwargs):
		return build_patents_workbook(
			subjects or [self.ms, self.pd],
			attribution_for_site(self.site),
			options_summary="test",
			**kwargs,
		)

	def export(self, **kwargs):
		kwargs.setdefault("site", str(self.site.pk))
		fd, path = tempfile.mkstemp(suffix=".xlsx")
		os.close(fd)
		self.addCleanup(os.unlink, path)
		call_command("export_patents_xlsx", output=path, **kwargs)
		return openpyxl.load_workbook(path)

	def ms_rows(self, wb=None):
		wb = wb or self.build()
		return {r["patent_id"]: r for r in rows_by_header(wb["Multiple Sclerosis"])}


class CommandTests(PatentExportTestCase):
	def test_sheet_order_and_names(self):
		wb = self.export(subjects=f"{self.ms.pk},{self.pd.pk}")
		self.assertEqual(wb.sheetnames, SHEET_ORDER)

	def test_all_subjects_and_team_filter(self):
		other_team = Team.objects.create(name="Other", organization=self.org, slug="other")
		Subject.objects.create(subject_name="Other subject", subject_slug="o", team=other_team)
		wb = self.export(all_subjects=True, team=self.team.pk)
		self.assertNotIn("Other subject", wb.sheetnames)
		self.assertIn("Parkinson", wb.sheetnames)

	def test_one_row_per_family_and_family_in_two_subjects_on_both_sheets(self):
		wb = self.export(subjects=f"{self.ms.pk},{self.pd.pk}")
		ms_ids = {r["patent_id"] for r in rows_by_header(wb["Multiple Sclerosis"])}
		pd_ids = {r["patent_id"] for r in rows_by_header(wb["Parkinson"])}
		self.assertEqual(ms_ids, {self.p1.pk, self.p2.pk, self.p3.pk})
		self.assertEqual(pd_ids, {self.p2.pk})

	def test_subject_without_patents_gets_a_note(self):
		empty = Subject.objects.create(subject_name="Empty", subject_slug="e", team=self.team)
		wb = self.export(subjects=str(empty.pk))
		self.assertEqual(wb["Empty"]["A2"].value, "No patents found for this subject.")

	def test_validation(self):
		with self.assertRaises(CommandError):
			call_command("export_patents_xlsx", site="a.example", output="/tmp/noop.xlsx")
		with self.assertRaises(CommandError):
			call_command("export_patents_xlsx", subjects="99999", site="a.example", output="/tmp/noop.xlsx")
		with self.assertRaises(CommandError):
			call_command("export_patents_xlsx", subjects="abc", site="a.example", output="/tmp/noop.xlsx")
		with self.assertRaises(CommandError):
			call_command("export_patents_xlsx", subjects=str(self.ms.pk), site="nope.example", output="/tmp/noop.xlsx")
		with self.assertRaises(CommandError):
			call_command("export_patents_xlsx", subjects=str(self.ms.pk), output="/tmp/noop.xlsx")
		other_team = Team.objects.create(name="Other", organization=self.org, slug="other")
		with self.assertRaises(CommandError):
			call_command(
				"export_patents_xlsx", subjects=str(self.ms.pk), team=other_team.pk,
				site="a.example", output="/tmp/noop.xlsx",
			)

	def test_site_by_domain(self):
		wb = self.export(subjects=str(self.ms.pk), site="a.example")
		self.assertEqual(about_values(wb)["Published by"], "Project A")


class SubjectSheetTests(PatentExportTestCase):
	def test_ordered_by_priority_date_with_null_last(self):
		wb = self.build()
		ids = [r["patent_id"] for r in rows_by_header(wb["Multiple Sclerosis"])]
		self.assertEqual(ids, [self.p2.pk, self.p1.pk, self.p3.pk])

	def test_identity_and_dates(self):
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(row["family_id"], "100")
		self.assertEqual(row["representative_publication"], "EP4123456A1")
		self.assertEqual(as_date(row["earliest_priority_date"]), D(2022, 5, 1))
		self.assertIsNotNone(row["discovery_date"])
		self.assertTrue(row["has_grant"])
		self.assertIn("espacenet", row["links_json"])

	def test_internal_fields_excluded(self):
		headers = headers_of(self.build()["Multiple Sclerosis"])
		for name in ("utitle", "usummary", "family_next_check", "family_attempts"):
			self.assertNotIn(name, headers)

	def test_long_claims_are_cut_off(self):
		Patents.objects.filter(pk=self.p1.pk).update(claims="x" * 40000)
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(len(row["claims"]), 32767)
		self.assertTrue(row["claims"].endswith("…"))
		self.assertEqual(row["claims_length"], 40000)
		self.assertTrue(row["claims_truncated"])

	def test_long_claims_starting_with_a_trigger_stay_within_the_cell_limit(self):
		# WO claims often open with a page number such as "- 27 -"
		Patents.objects.filter(pk=self.p1.pk).update(claims="- 27 -" + "x" * 40000)
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(len(row["claims"]), 32767)
		self.assertTrue(row["claims"].startswith("'- 27 -"))
		self.assertTrue(row["claims_truncated"])

	def test_claims_one_short_of_the_limit_with_a_trigger_are_flagged(self):
		Patents.objects.filter(pk=self.p1.pk).update(claims="-" + "x" * 32766)
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(len(row["claims"]), 32767)
		self.assertTrue(row["claims_truncated"])

	def test_control_characters_are_dropped(self):
		Patents.objects.filter(pk=self.p1.pk).update(
			title="Anti\x0bbody", claims="1. An\x01 antibody."
		)
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(row["title"], "Antibody")
		self.assertEqual(row["claims"], "1. An antibody.")

	def test_short_claims_are_whole(self):
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(row["claims"], "1. An antibody.")
		self.assertEqual(row["claims_length"], len("1. An antibody."))
		self.assertFalse(row["claims_truncated"])

	def test_applicant_columns(self):
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(
			row["applicants"],
			"NOVARTIS AG → Novartis (sponsor %d; industry; curated); "
			"SMITH, John (individual); ACME HOLDINGS IP (no sponsor match)" % self.novartis.pk,
		)
		self.assertEqual(row["applicants_raw"], "NOVARTIS AG; SMITH, John; ACME HOLDINGS IP")
		self.assertEqual(row["applicants_individual"], "SMITH, John")
		self.assertEqual(row["applicant_count"], 3)
		self.assertTrue(row["has_individual_applicant"])
		self.assertTrue(row["has_unmatched_applicant"])
		self.assertEqual(row["applicant_sponsor_ids"], str(self.novartis.pk))
		self.assertEqual(row["applicant_sponsor_slugs"], "novartis")
		self.assertEqual(row["applicant_sponsor_names"], "Novartis")
		self.assertEqual(row["applicant_sponsor_types"], "industry")
		self.assertEqual(row["applicant_sponsor_type_sources"], "curated")

	def test_matched_sponsor_without_type_keeps_positions(self):
		bare = Sponsor.objects.create(name="Bare Co", slug="bare-co")
		PatentApplicant.objects.create(patent=self.p3, sponsor=bare, raw_name="BARE CO", sequence=1)
		row = self.ms_rows()[self.p3.pk]
		self.assertIn("BARE CO → Bare Co (sponsor %d; type unknown; no source)" % bare.pk, row["applicants"])
		self.assertIn(row["applicant_sponsor_types"], (None, ""))

	def test_inventors_and_classification(self):
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(row["inventors"], "Jane Doe; John Roe")
		self.assertEqual(row["inventor_count"], 2)
		self.assertEqual(row["cpc_classes"], "A61K39/395; A61P25/00")
		self.assertEqual(row["cpc_count"], 2)

	def test_non_string_inventor_does_not_crash(self):
		Patents.objects.filter(pk=self.p1.pk).update(inventors=["Jane", {"name": "x"}, 5])
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(row["inventor_count"], 3)

	def test_publication_summary(self):
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(
			row["publication_numbers"], "WO2022123456A1; EP4123456A1; EP4123456B1"
		)
		self.assertEqual(row["publication_count"], 3)
		self.assertEqual(row["publication_offices"], "WO; EP")
		self.assertEqual(row["granted_publications"], "EP4123456B1")
		self.assertEqual(row["application_numbers"], "PCT/EP2022/1; EP22000001")
		self.assertEqual(as_date(row["earliest_application_date"]), D(2022, 5, 1))

	def test_relations_and_categories(self):
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(row["subjects"], "Multiple Sclerosis")
		self.assertEqual(row["teams"], "Team A")
		self.assertEqual(row["sources"], "EPO | MS")
		self.assertEqual(
			sorted(row["categories"].split("; ")),
			["Anti-CD20 (manual)", "Ocrelizumab (auto)"],
		)

	def test_trial_link_columns(self):
		row = self.ms_rows()[self.p1.pk]
		self.assertEqual(row["trial_link_count"], 4)
		self.assertEqual(
			row["trial_links_summary"], "category: 1 during, 1 after; subject: 1 during; suppressed: 1"
		)

	def test_no_links_summary_is_empty(self):
		row = self.ms_rows()[self.p2.pk]
		self.assertEqual(row["trial_link_count"], 0)
		self.assertIn(row["trial_links_summary"], (None, ""))


class TrialLinksSheetTests(PatentExportTestCase):
	def rows(self):
		return rows_by_header(self.build()["Trial links"])

	def test_all_links_including_suppressed(self):
		rows = self.rows()
		self.assertEqual(len(rows), 4)
		suppressed = [r for r in rows if r["suppressed"]]
		self.assertEqual([r["trial_title"] for r in suppressed], ["Dismissed trial"])

	def test_sort_order_basis_then_suppressed(self):
		rows = self.rows()
		self.assertEqual(
			[(r["basis"], r["suppressed"]) for r in rows],
			[("category", False), ("category", False), ("subject", False), ("subject", True)],
		)

	def test_sort_by_priority_date_within_group(self):
		other = self.make_patent("400", "Older", D(2019, 1, 1), subjects=[self.ms])
		self.make_link(other, self.trial_closed, "during", "category")
		rows = [r for r in self.rows() if r["basis"] == "category"]
		self.assertEqual(rows[-1]["patent_id"], other.pk)

	def test_both_bases_present(self):
		self.assertEqual({r["basis"] for r in self.rows()}, {"category", "subject"})

	def test_window_columns_match_detector(self):
		rows = {r["trial_id"]: r for r in self.rows()}
		for trial in (self.trial_open, self.trial_closed):
			start, end = trial_window(trial)
			row = rows[trial.pk]
			self.assertEqual(as_date(row["trial_window_start"]), start)
			if end is None:
				self.assertIn(row["trial_window_end"], (None, ""))
			else:
				self.assertEqual(as_date(row["trial_window_end"]), end)

	def test_shared_overlap_and_trial_fields(self):
		rows = {r["trial_id"]: r for r in self.rows()}
		row = rows[self.trial_open.pk]
		self.assertEqual(row["shared_subjects"], "Multiple Sclerosis")
		self.assertEqual(row["shared_categories"], "Ocrelizumab")
		self.assertEqual(row["sponsor_name"], "Novartis")
		self.assertEqual(row["primary_outcome"], "Relapse rate")
		self.assertEqual(row["patent_title"], "Anti-CD20 antibody")
		self.assertEqual(rows[self.trial_closed.pk]["days_after_completion"], 120)

	def test_link_to_patent_outside_export_is_left_out(self):
		wb = self.build(subjects=[self.pd])
		self.assertEqual(wb["Trial links"]["A2"].value, "No trial links found.")


class PublicationsSheetTests(PatentExportTestCase):
	def test_one_row_per_member(self):
		wb = self.build()
		rows = rows_by_header(wb["Publications"])
		self.assertEqual(len(rows), 5)
		p1_rows = [r for r in rows if r["patent_id"] == self.p1.pk]
		self.assertEqual(
			[r["publication_number"] for r in p1_rows],
			["WO2022123456A1", "EP4123456A1", "EP4123456B1"],
		)

	def test_representative_and_grant_flags(self):
		rows = {r["publication_number"]: r for r in rows_by_header(self.build()["Publications"])}
		self.assertTrue(rows["EP4123456A1"]["is_representative"])
		self.assertFalse(rows["EP4123456B1"]["is_representative"])
		self.assertTrue(rows["EP4123456B1"]["is_grant_kind"])
		self.assertFalse(rows["EP4123456A1"]["is_grant_kind"])
		self.assertFalse(rows["WO2022123456A1"]["is_grant_kind"])
		self.assertEqual(rows["EP4123456A1"]["family_id"], "100")
		self.assertIn("CC=EP&NR=4123456A1", rows["EP4123456A1"]["espacenet_link"])
		self.assertEqual(rows["EP4123456A1"]["sources"], "epo_ops")


class CategoriesAboutAndGlossaryTests(PatentExportTestCase):
	def test_categories_sheet_uses_patent_settings(self):
		wb = self.build()
		ws = wb["Categories"]
		rows = rows_by_header(ws, header_row=4)
		by_name = {r["Category"]: r for r in rows}
		self.assertIn("Min score (patents)", rows[0])
		self.assertIn("Field weights (patents)", rows[0])
		self.assertEqual(by_name["Ocrelizumab"]["Patents (this subject)"], 1)
		self.assertEqual(by_name["Ocrelizumab"]["Min score (patents)"], 3)
		self.assertIn("title:3", by_name["Ocrelizumab"]["Field weights (patents)"])

	def test_about_sheet_attribution_and_terms_always_present(self):
		self.custom.delete()
		wb = self.build()
		about = about_values(wb)
		self.assertIn("Open Patent Services", about["Patent data"])
		self.assertEqual(about["Terms of use"], "For research use; do not republish the patent text.")
		self.assertIn("Patents export", about["How to cite"])
		self.assertEqual(about["Multiple Sclerosis"], '3 patent families for the research subject "Multiple Sclerosis"')

	def test_about_licence_label(self):
		self.custom.data_license = "CC BY 4.0"
		self.custom.save()
		about = about_values(self.build())
		self.assertEqual(about["Data licence (Gregory's processing)"], "CC BY 4.0")
		self.assertNotIn("Data licence", about)

	def test_about_lists_fixed_sheets(self):
		about = about_values(self.build())
		for name in ("Trial links", "Publications", "Categories", "Glossary", "Data source"):
			self.assertIn(name, about)

	def test_glossary_covers_every_column_of_every_data_sheet(self):
		wb = self.build()
		glossary = {r["Field"]: r for r in rows_by_header(wb["Glossary"])}
		for sheet in ("Multiple Sclerosis", "Parkinson", "Trial links", "Publications"):
			for header in headers_of(wb[sheet]):
				self.assertIn(header, glossary, f"{header} missing from glossary")
				self.assertTrue(
					glossary[header]["Description"], f"{header} has no description"
				)
				self.assertTrue(glossary[header]["Label"])
		self.assertIn("Trial links", glossary["patent_id"]["Sheet"])
		self.assertIn("Publications", glossary["patent_id"]["Sheet"])

	def test_glossary_covers_declared_columns(self):
		for col in TRIAL_LINK_COLS + PUBLICATION_COLS + patents_xlsx._PLACED_COLS:
			self.assertIn(col, patents_xlsx.PATENT_GLOSSARY, col)

	def test_data_source_sheet_covers_all_topics(self):
		wb = self.build()
		text = "\n".join(
			str(c.value) for row in wb["Data source"].iter_rows() for c in row if c.value
		)
		self.assertEqual(len(patents_xlsx.DATA_SOURCE_SECTIONS), 17)
		for heading, _, _ in patents_xlsx.DATA_SOURCE_SECTIONS:
			self.assertIn(heading, text)
		self.assertIn("EP B1", text)


class FormulaInjectionTests(PatentExportTestCase):
	def test_untrusted_text_is_written_as_text(self):
		for i, trigger in enumerate(("=", "+", "-", "@")):
			patent = self.make_patent(
				f"9{i}",
				f"{trigger}title{i}",
				D(2024, 1, 1 + i),
				subjects=[self.pd],
				claims=f"{trigger}claims{i}",
				inventors=[f"{trigger}inventor{i}"],
			)
			PatentApplicant.objects.create(
				patent=patent, raw_name=f"{trigger}applicant{i}", sequence=1, is_individual=True
			)
		wb = self.build()
		ws = wb["Parkinson"]
		headers = headers_of(ws)
		checked = 0
		for row in ws.iter_rows(min_row=2):
			cells = dict(zip(headers, row))
			if not str(cells["family_id"].value).startswith("9"):
				continue
			for name in ("title", "claims", "inventors", "applicants_raw", "applicants"):
				cell = cells[name]
				self.assertTrue(cell.value.startswith("'"), (name, cell.value))
				self.assertEqual(cell.data_type, "s")
			checked += 1
		self.assertEqual(checked, 4)


class QueryCountTests(PatentExportTestCase):
	def count_queries(self):
		with CaptureQueriesContext(connection) as ctx:
			self.build()
		return len(ctx)

	def test_query_count_does_not_grow_with_patents(self):
		self.count_queries()  # warm caches (Site framework etc.)
		before = self.count_queries()
		for i in range(9):
			patent = self.make_patent(f"5{i}", f"Extra {i}", D(2020, 1, 1 + i), subjects=[self.ms, self.pd])
			self.add_pub(patent, "EP", f"70000{i}", "A1", D(2021, 1, 1))
			self.add_pub(patent, "WO", f"80000{i}", "A1", D(2021, 2, 1))
			sponsor = Sponsor.objects.create(name=f"Sponsor {i}", slug=f"sponsor-{i}", sponsor_type="industry")
			PatentApplicant.objects.create(patent=patent, sponsor=sponsor, raw_name=f"S{i}", sequence=1)
			PatentCategoryAssignment.objects.create(patents=patent, teamcategory=self.cat_auto, source="automatic")
			trial = self.make_trial(f"Extra trial {i}", "recruiting", D(2019, 1, 1), None)
			self.make_link(patent, trial, "during", "subject", sponsor=sponsor)
		after = self.count_queries()
		self.assertEqual(after, before)


class AdminExportTests(PatentExportTestCase):
	def setUp(self):
		super().setUp()
		# Organisation B: its names must never reach a staff member of A.
		self.org_b = Organization.objects.create(name="Org B", slug="org-b")
		self.team_b = Team.objects.create(name="BTEAMX", organization=self.org_b, slug="team-b")
		self.site_b = Site.objects.create(domain="b.example", name="Site B")
		OrganizationSite.objects.create(organization=self.org_b, site=self.site_b, is_default=True)
		self.subject_b = Subject.objects.create(subject_name="BSUBJECTX", subject_slug="bs", team=self.team_b)
		self.source_b = Sources.objects.create(name="BSOURCEX", source_for="patents", team=self.team_b)
		self.cat_b = TeamCategory.objects.create(team=self.team_b, category_name="BCATX", category_terms=["b"])
		self.cat_b.subjects.add(self.subject_b)
		self.p1.subjects.add(self.subject_b)
		self.p1.teams.add(self.team_b)
		self.p1.sources.add(self.source_b)
		PatentCategoryAssignment.objects.create(patents=self.p1, teamcategory=self.cat_b, source="manual")
		self.trial_b = self.make_trial("BTRIALX", "recruiting", D(2020, 1, 1), None, subject=self.subject_b)
		self.trial_b.sources.add(self.source_b)
		self.make_link(self.p1, self.trial_b, "during", "category")
		# Trials of A are visible to A through their source.
		for trial in (self.trial_closed, self.trial_open, self.trial_subject, self.trial_dismissed):
			trial.sources.add(self.source)

		self.staff = User.objects.create_user(username="staff-a", password="pw", is_staff=True)
		OrganizationUser.objects.create(user=self.staff, organization=self.org)
		self.staff.user_permissions.add(
			*Permission.objects.filter(codename__in=["view_patents", "view_patenttriallink"])
		)
		self.superuser = User.objects.create_superuser("root", "root@example.com", "pw")
		self.url = reverse("admin:gregory_patents_changelist")
		self.links_url = reverse("admin:gregory_patenttriallink_changelist")

	def login_client(self, user):
		client = Client()
		client.force_login(user)
		return client

	def download(self, user, pks, site, url=None, action="export_xlsx_action", **extra):
		client = self.login_client(user)
		data = {"action": action, "apply": "1", "site": site.pk, "_selected_action": pks}
		data.update(extra)
		response = client.post(url or self.url, data)
		self.assertEqual(response.status_code, 200, response.content[:300])
		return response, openpyxl.load_workbook(BytesIO(response.content))

	def all_text(self, wb):
		return "\n".join(
			str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value
		)

	def test_intermediate_page_offers_only_own_sites(self):
		response = self.login_client(self.staff).post(
			self.url, {"action": "export_xlsx_action", "_selected_action": [self.p1.pk]}
		)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, "a.example")
		self.assertNotContains(response, "b.example")
		self.assertContains(response, "Download XLSX")

	def test_superuser_sees_every_site(self):
		response = self.login_client(self.superuser).post(
			self.url, {"action": "export_xlsx_action", "_selected_action": [self.p1.pk]}
		)
		self.assertContains(response, "a.example")
		self.assertContains(response, "b.example")

	def test_staff_cannot_pick_other_organisations_site(self):
		response = self.login_client(self.staff).post(
			self.url,
			{"action": "export_xlsx_action", "apply": "1", "site": self.site_b.pk, "_selected_action": [self.p1.pk]},
		)
		self.assertNotEqual(
			response.get("Content-Disposition", ""), 'attachment; filename="x"'
		)
		self.assertNotIn("spreadsheetml", response.get("Content-Type", ""))

	def test_download_is_valid_xlsx(self):
		response, wb = self.download(self.staff, [self.p1.pk], self.site)
		self.assertEqual(
			response["Content-Type"],
			"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
		)
		self.assertRegex(response["Content-Disposition"], r'filename="patents_export_\d{8}\.xlsx"')
		self.assertIn("Multiple Sclerosis", wb.sheetnames)
		self.assertEqual(about_values(wb)["Published by"], "Project A")

	def test_staff_export_leaves_out_other_organisations_data(self):
		_, wb = self.download(self.staff, [self.p1.pk], self.site)
		text = self.all_text(wb)
		for secret in ("BTEAMX", "BSOURCEX", "BCATX", "BSUBJECTX", "BTRIALX"):
			self.assertNotIn(secret, text)
		self.assertNotIn("BSUBJECTX", wb.sheetnames)
		links = rows_by_header(wb["Trial links"])
		self.assertNotIn(self.trial_b.pk, {r["trial_id"] for r in links})
		self.assertIn(self.trial_open.pk, {r["trial_id"] for r in links})
		row = rows_by_header(wb["Multiple Sclerosis"])[0]
		self.assertEqual(row["subjects"], "Multiple Sclerosis")
		self.assertEqual(row["teams"], "Team A")
		self.assertEqual(row["sources"], "EPO | MS")
		self.assertEqual(row["trial_link_count"], 4)

	def test_superuser_export_shows_everything(self):
		_, wb = self.download(self.superuser, [self.p1.pk], self.site)
		text = self.all_text(wb)
		for name in ("BTEAMX", "BSOURCEX", "BCATX", "BSUBJECTX", "BTRIALX"):
			self.assertIn(name, text)

	def test_patent_with_no_subject_lands_on_no_subject_sheet(self):
		orphan = self.make_patent("600", "Orphan patent", D(2020, 1, 1), subjects=[])
		_, wb = self.download(self.staff, [orphan.pk, self.p1.pk], self.site)
		self.assertIn("No subject", wb.sheetnames)
		self.assertEqual(
			[r["patent_id"] for r in rows_by_header(wb["No subject"])], [orphan.pk]
		)
		self.assertIn("No subject", about_values(wb))
		# the patent only visible through another organisation's subject also lands there
		hidden = self.make_patent("700", "Hidden", D(2020, 1, 1), subjects=[self.subject_b])
		_, wb = self.download(self.staff, [hidden.pk], self.site)
		self.assertEqual(
			[r["patent_id"] for r in rows_by_header(wb["No subject"])], [hidden.pk]
		)
		self.assertNotIn("BSUBJECTX", self.all_text(wb))

	def test_no_subject_sheet_absent_when_not_needed(self):
		_, wb = self.download(self.staff, [self.p1.pk], self.site)
		self.assertNotIn("No subject", wb.sheetnames)

	def test_select_across_with_changelist_filter(self):
		url = self.url + "?has_grant__exact=1"
		client = self.login_client(self.staff)
		first = client.post(url, {"action": "export_xlsx_action", "select_across": "1", "index": "0", "_selected_action": [self.p2.pk]})
		self.assertEqual(first.status_code, 200)
		self.assertContains(first, 'name="select_across" value="1"')
		self.assertContains(first, 'name="_selected_action" value="0"')
		self.assertContains(first, "<strong>1 patent family</strong>")
		response = client.post(
			url,
			{
				"action": "export_xlsx_action",
				"select_across": "1",
				"index": "0",
				"_selected_action": "0",
				"apply": "1",
				"site": self.site.pk,
			},
		)
		wb = openpyxl.load_workbook(BytesIO(response.content))
		ids = {r["patent_id"] for r in rows_by_header(wb["Multiple Sclerosis"])}
		self.assertEqual(ids, {self.p1.pk})

	def test_select_across_without_filter_exports_everything_visible(self):
		client = self.login_client(self.staff)
		response = client.post(
			self.url,
			{"action": "export_xlsx_action", "select_across": "1", "index": "0", "_selected_action": "0", "apply": "1", "site": self.site.pk},
		)
		wb = openpyxl.load_workbook(BytesIO(response.content))
		ids = {r["patent_id"] for r in rows_by_header(wb["Multiple Sclerosis"])}
		self.assertEqual(ids, {self.p1.pk, self.p2.pk, self.p3.pk})

	def test_trial_link_action_exports_distinct_patents(self):
		pks = [self.link_cat_during.pk, self.link_cat_after.pk, self.link_subj.pk]
		_, wb = self.download(
			self.staff, pks, self.site, url=self.links_url, action="export_patents_xlsx_action"
		)
		rows = rows_by_header(wb["Multiple Sclerosis"])
		self.assertEqual([r["patent_id"] for r in rows], [self.p1.pk])
		self.assertEqual(len(rows_by_header(wb["Publications"])), 3)

	def test_trial_link_action_intermediate_page(self):
		response = self.login_client(self.staff).post(
			self.links_url,
			{"action": "export_patents_xlsx_action", "_selected_action": [self.link_cat_during.pk, self.link_cat_after.pk]},
		)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, "<strong>1 patent family</strong>")
		self.assertContains(response, 'value="export_patents_xlsx_action"')


class ScopeBuilderTests(PatentExportTestCase):
	def test_scope_hides_other_organisation_categories_on_categories_sheet(self):
		org_b = Organization.objects.create(name="Org B", slug="org-b2")
		team_b = Team.objects.create(name="B", organization=org_b, slug="team-b2")
		cat_b = TeamCategory.objects.create(team=team_b, category_name="BCAT", category_terms=["b"])
		cat_b.subjects.add(self.ms)
		unscoped = self.build()
		scoped = self.build(scope=ExportScope(org_ids=[self.org.pk]))
		names = lambda wb: {r["Category"] for r in rows_by_header(wb["Categories"], header_row=4)}
		self.assertIn("BCAT", names(unscoped))
		self.assertNotIn("BCAT", names(scoped))
