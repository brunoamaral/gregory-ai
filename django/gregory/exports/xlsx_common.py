"""Helpers shared by the trials and patents XLSX exports."""

import json
import re
from dataclasses import dataclass
from datetime import date, datetime, timezone as dt_timezone

from django.contrib.sites.models import Site
from django.core.management.base import CommandError
from django.db.models import Count
from django.db.models.functions import Lower
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from gregory.models import OrganizationSite, Patents, TeamCategory, Trials
from sitesettings.models import CustomSetting


def _sanitise_sheet_name(name, used):
	"""Return a valid, unique Excel sheet name (≤31 chars, no illegal characters)."""
	name = re.sub(r"[\\/*?\[\]:]", "_", name)[:31]
	base, suffix = name, 1
	while name in used:
		suffix += 1
		name = base[:28] + f"_{suffix}"
	used.add(name)
	return name


def _cell_value(value):
	"""Convert a Python value to an Excel-safe type (None → empty string)."""
	if value is None:
		return ""
	if isinstance(value, datetime):
		if value.tzinfo is not None:
			value = value.astimezone(dt_timezone.utc).replace(tzinfo=None)
		return value
	if isinstance(value, date):
		return value
	if isinstance(value, dict):
		return json.dumps(value, ensure_ascii=False)
	if isinstance(value, list):
		# Lists of scalars (e.g. regions_normalized) render as a "; "-joined string;
		# lists containing dicts/lists fall back to JSON for fidelity.
		if all(not isinstance(v, (dict, list)) for v in value):
			return "; ".join("" if v is None else str(v) for v in value)
		return json.dumps(value, ensure_ascii=False)
	if isinstance(value, bool):
		return value
	if isinstance(value, (int, float)):
		return value
	return str(value)


def _apply_header(ws, columns, row=1):
	"""Write a bold, coloured header row and return the cell count."""
	hdr_font = Font(bold=True, color="FFFFFF")
	hdr_fill = PatternFill(fill_type="solid", fgColor="2F4F8F")
	hdr_align = Alignment(vertical="center")
	for col_idx, name in enumerate(columns, 1):
		cell = ws.cell(row=row, column=col_idx, value=name)
		cell.font = hdr_font
		cell.fill = hdr_fill
		cell.alignment = hdr_align
	return len(columns)


def _excel_text(value, limit=32767):
	"""Truncate a string to Excel's per-cell character limit."""
	if not value:
		return value
	text = str(value)
	if len(text) > limit:
		text = text[: limit - 1] + "…"
	return text


_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@")


def _write_safe_text_cell(ws, row, col, value, wrap=False):
	"""Write a long/user-authored text value, truncated and defused against formula injection.

	A leading apostrophe is the standard mitigation (OWASP CSV injection guidance):
	it forces the cell to render as literal text in Excel/LibreOffice/Sheets even
	when the value starts with =, +, -, or @. Relying on openpyxl's data_type alone
	is not enough — several of those trigger characters (+, -, @) are never
	auto-detected as formulas by openpyxl in the first place, so its cell.data_type
	would stay "s" regardless, while spreadsheet applications still treat a leading
	=, +, -, or @ as a formula cue on their own.
	"""
	text = _excel_text(value)
	if isinstance(text, str) and text.startswith(_FORMULA_TRIGGER_CHARS):
		text = "'" + text
	cell = ws.cell(row=row, column=col, value=text)
	cell.data_type = "s"
	if wrap:
		cell.alignment = Alignment(wrap_text=True)
	return cell


GENERATED_BY_PROSE = "GregoryAI — https://github.com/brunoamaral/gregory-ai"


@dataclass
class SiteAttribution:
	"""Which Site (and its CustomSetting) an export is attributed to."""

	site: object
	custom_setting: object = None


def _resolve_explicit_site(value):
	"""Resolve --site by numeric ID or domain (case-insensitive). Raises CommandError if not found."""
	value = value.strip()
	if value.isdigit():
		site = Site.objects.filter(pk=int(value)).first()
	else:
		site = Site.objects.filter(domain__iexact=value).first()
	if site is None:
		listing = ", ".join(f"{s.pk} ({s.domain})" for s in Site.objects.order_by("pk"))
		raise CommandError(f"Site not found: {value!r}. Valid sites: {listing}")
	return site


def _resolve_site_attribution(explicit_site_value):
	"""Resolve the SiteAttribution for an export from the required --site value.

	The export used to infer a site per subject from Team.site, falling back to
	the subject's organisation's default site -- removed along with Team.site
	(SITE-API-VISIBILITY-SPEC.md, "Removing Team.site"). --site is now the only
	source, which also removes the case where one export mixed rows resolved
	against different sites.
	"""
	site = _resolve_explicit_site(explicit_site_value)
	custom_setting = (
		CustomSetting.objects.filter(site=site).order_by("setting_id").first()
	)
	return SiteAttribution(site=site, custom_setting=custom_setting)


def attribution_for_site(site):
	"""SiteAttribution for an already-resolved Site (used by the admin actions)."""
	custom_setting = CustomSetting.objects.filter(site=site).order_by("setting_id").first()
	return SiteAttribution(site=site, custom_setting=custom_setting)


def _category_rows(subjects, *, content_type, model, category_qs=None, count_qs=None):
	"""Return one row per (subject, category) pair, ordered by subject then category name.

	Categories with no subject never appear here: filtering by `subjects=subject`
	only matches categories reachable from an exported subject.

	``content_type`` is "trial" or "patent". ``category_qs`` limits which categories
	may appear (the admin passes the caller's organisations' categories) and
	``count_qs`` is the base queryset the per-category counts are taken from.
	"""
	min_score_field = f"match_min_score_{content_type}s"
	if category_qs is None:
		category_qs = TeamCategory.objects.all()
	if count_qs is None:
		count_qs = model.objects.all()
	rows = []
	for subject in subjects:
		cats = list(
			category_qs.filter(subjects=subject)
			.select_related("team")
			.order_by(Lower("category_name"))
		)
		if not cats:
			continue
		counts = {
			r["team_categories"]: r["n"]
			for r in (
				count_qs.filter(
					subjects=subject, team_categories__in=[c.pk for c in cats]
				)
				.values("team_categories")
				.annotate(n=Count("pk", distinct=True))
			)
		}
		for cat in cats:
			weights = cat.get_scored_fields(content_type)
			weights_str = "; ".join(f"{f}:{w}" for f, w in weights.items() if w)
			rows.append(
				(
					subject.subject_name,
					cat.team.name,
					cat.category_name,
					cat.category_slug or "",
					cat.category_description or "",
					"; ".join(cat.category_terms or []),
					len(cat.category_terms or []),
					"; ".join(cat.category_exclude_terms or []),
					cat.get_category_type_display(),
					cat.get_modality_display() if cat.modality else "",
					cat.get_match_scope_display(),
					getattr(cat, min_score_field),
					weights_str,
					counts.get(cat.pk, 0),
					_cell_value(cat.last_synced_at),
				)
			)
	return rows


def _build_categories_sheet(
	wb, subjects, *, content_type, columns, prose, category_qs=None, count_qs=None
):
	"""Add a Categories sheet — one row per (subject, category) pair."""
	model = Patents if content_type == "patent" else Trials
	ws = wb.create_sheet(title="Categories")
	section_font = Font(bold=True, size=13)

	ws.cell(row=1, column=1, value="Categories and their search terms").font = section_font
	prose_cell = ws.cell(row=2, column=1, value=prose)
	prose_cell.alignment = Alignment(wrap_text=True)
	last_col_letter = get_column_letter(len(columns))
	ws.merge_cells(f"A2:{last_col_letter}2")
	ws.row_dimensions[2].height = 90

	_apply_header(ws, columns, row=4)
	ws.freeze_panes = "A5"
	ws.auto_filter.ref = f"A4:{last_col_letter}4"

	rows = _category_rows(
		subjects,
		content_type=content_type,
		model=model,
		category_qs=category_qs,
		count_qs=count_qs,
	)
	desc_col = columns.index("Description") + 1
	terms_col = columns.index("Search terms") + 1
	exclude_col = columns.index("Excluded terms") + 1

	if not rows:
		ws.cell(row=5, column=1, value="No categories found.")
	else:
		for row_idx, row_data in enumerate(rows, 5):
			for col_idx, value in enumerate(row_data, 1):
				if col_idx in (desc_col, terms_col, exclude_col):
					_write_safe_text_cell(ws, row_idx, col_idx, value, wrap=True)
				else:
					ws.cell(row=row_idx, column=col_idx, value=value)

	wide_widths = {"Description": 65, "Search terms": 65, "Excluded terms": 40}
	for col_idx, name in enumerate(columns, 1):
		letter = get_column_letter(col_idx)
		if name.startswith("Field weights"):
			ws.column_dimensions[letter].width = 40
		else:
			ws.column_dimensions[letter].width = wide_widths.get(
				name, max(12, min(32, len(name) + 4))
			)


def _build_about_sheet(
	ws,
	attribution,
	sheet_entries,
	options_summary,
	*,
	item_noun,
	citation_label,
	fixed_rows,
	extra_source_rows=(),
	licence_label="Data licence",
):
	"""Populate the 'About this file' sheet: Source, This file, What's in this workbook.

	``item_noun(count)`` renders e.g. "3 clinical trial(s)"; ``citation_label`` is the
	noun in the default citation; ``fixed_rows`` are (sheet, description) pairs for the
	non-subject sheets; ``extra_source_rows`` are (label, value) pairs written at the end
	of the Source section whatever the site's settings.
	"""
	section_fill = PatternFill(fill_type="solid", fgColor="2F4F8F")
	section_font = Font(bold=True, size=13, color="FFFFFF")
	label_font = Font(bold=True)

	ws.column_dimensions["A"].width = 28
	ws.column_dimensions["B"].width = 90

	row = 1

	def write_section_header(text):
		nonlocal row
		cell = ws.cell(row=row, column=1, value=text)
		cell.font = section_font
		cell.fill = section_fill
		ws.cell(row=row, column=2).fill = section_fill
		ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
		row += 1

	def write_row(label, value, wrap=True):
		nonlocal row
		if not value:
			return
		# Most labels are our own fixed strings, but Section 3's are subject
		# names, so route every label through the same defusing as the value.
		_write_safe_text_cell(ws, row, 1, label, wrap=False)
		ws.cell(row=row, column=1).font = label_font
		_write_safe_text_cell(ws, row, 2, value, wrap=wrap)
		row += 1

	site = attribution.site
	cs = attribution.custom_setting

	write_section_header("Source")
	title = (cs.title if cs and cs.title else None) or site.name or site.domain
	write_row("Published by", title)

	org_site = (
		OrganizationSite.objects.filter(site=site).select_related("organization").first()
	)
	org_name = org_site.organization.name if org_site else ""
	if org_name and org_name != title:
		write_row("Organisation", org_name)

	if cs:
		write_row("About this project", cs.description)
		write_row("Website", cs.website_url or f"https://{site.domain}")
		write_row("About page", cs.about_url)
		write_row("Contact", cs.contact_url)
		write_row("Contact email", cs.contact_email or cs.admin_email)
		if cs.api_domain:
			write_row("API", f"https://{cs.api_domain}")
		write_row("Source code", cs.github_url)
		social = "; ".join(
			filter(
				None,
				[
					f"Mastodon: {cs.mastodon_url}" if cs.mastodon_url else "",
					f"Bluesky: {cs.bluesky_url}" if cs.bluesky_url else "",
				],
			)
		)
		write_row("Mastodon / Bluesky", social)
		write_row("Privacy policy", cs.privacy_policy_url)
	else:
		write_row("Website", f"https://{site.domain}")

	for label, value in extra_source_rows:
		write_row(label, value)

	write_section_header("This file")
	generated = datetime.now(dt_timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
	write_row("Generated", generated)
	write_row("Generated by", GENERATED_BY_PROSE)
	write_row("Export options", options_summary)

	data_license = cs.data_license if cs else ""
	data_license_url = cs.data_license_url if cs else ""
	write_row(licence_label, data_license)
	write_row("Licence URL", data_license_url)

	citation = cs.citation if cs else ""
	if not citation:
		title_for_citation = (
			(cs.title if cs and cs.title else None) or site.name or site.domain
		)
		website_for_citation = (cs.website_url if cs else "") or f"https://{site.domain}"
		date_str = date.today().strftime("%Y-%m-%d")
		citation = f"{title_for_citation}. {citation_label}, {date_str}."
		if website_for_citation:
			citation += f" {website_for_citation}"
	write_row("How to cite", citation)

	write_section_header("What's in this workbook")
	for sheet_name, subject, count in sheet_entries:
		value = f'{item_noun(count)} for the research subject "{subject.subject_name}"'
		if subject.description:
			value += f" — {subject.description}"
		write_row(sheet_name, value)
	for label, value in fixed_rows:
		write_row(label, value)
